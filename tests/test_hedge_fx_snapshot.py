"""Synthetic-only source/date/unit tests for private Hedge/FX imports."""
from datetime import date
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from build_hedge_fx_snapshot import (EXPECTED_HEADERS, RATE_FIELDS, build,
                                     build_histories, model_input, read_history)
from build_info_snapshot import HEADERS, REPO


def values(**updates):
    return {**dict.fromkeys(EXPECTED_HEADERS), **updates}


class TestHedgeFxReader(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'synthetic.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows, transform=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'DailyRate'
        sheet.cell(8, 1, 'DATES')
        sheet.cell(9, 1, '일자')
        for column, labels in {**HEADERS, **EXPECTED_HEADERS}.items():
            index = column_index_from_string(column)
            sheet.cell(8, index, labels[0])
            sheet.cell(9, index, labels[1])
        for n, (observed, data) in enumerate(rows, 10):
            sheet.cell(n, 1, observed)
            for column, value in data.items():
                sheet.cell(n, column_index_from_string(column), value)
        if transform:
            transform(sheet)
        book.save(self.path)
        book.close()
        return self.path

    def test_cutoff_sorted_dates_signs_missing_and_weekends(self):
        path = self.workbook([
            (date(2026, 10, 15), {'V': 1300, 'HW': -1, 'HX': 0, 'HY': 1, 'B': 2, 'AE': 3}),
            (date(2026, 10, 14), {'V': 1290, 'HW': -2, 'HX': None, 'HY': -3}),
            (date(2026, 10, 15), {'V': 1300, 'HW': -1, 'HX': 0, 'HY': 1, 'B': 2, 'AE': 3}),
            (date(2026, 10, 16), {'V': 99999, 'HW': 99999}),
            (date(2026, 10, 11), {'V': 88888}),
            (date(2014, 12, 31), {'V': 77777}),
        ])
        rows, audit = read_history(path, date(2026, 10, 15))
        self.assertEqual([d.isoformat() for d, _ in rows], ['2026-10-14', '2026-10-15'])
        self.assertEqual(audit['identicalDuplicateRows'], 1)
        self.assertEqual(audit['futureRowsExcluded'], 1)
        self.assertEqual(audit['weekendRowsExcluded'], 1)
        self.assertEqual(audit['beforeStartRowsExcluded'], 1)
        result = build_histories(rows, audit)
        usd = result['pairs'][0]
        self.assertEqual(usd['history'][-1]['hedgeMean'], 0)
        self.assertIsNone(usd['history'][0]['hedgeMean'])
        self.assertEqual(usd['history'][0]['hedge3m'], -2)
        self.assertEqual(result['usdRates'][-1]['kr3m'] - result['usdRates'][-1]['us3m'], -1)
        self.assertEqual(result['unit'], 'source_value')
        self.assertFalse(result['unitVerified'])
        self.assertEqual(result['quality']['pairs']['USDKRW']['incompleteHedgeRows'], 1)

    def test_headers_fail_closed_when_source_definition_or_location_changes(self):
        for transform in [lambda s: s.__setitem__('AE9', 'Yield'),
                          lambda s: s.__setitem__('B8', '한국_6M'),
                          lambda s: s.__setitem__('HW9', '연율(%)'),
                          lambda s: (s.__setitem__('C8', 'UST3M'), s.__setitem__('C9', 'MID_Close'))]:
            path = self.workbook([(date(2026, 10, 15), {'V': 1300})], transform)
            with self.assertRaisesRegex(ValueError, 'Changed or duplicate'):
                read_history(path, date(2026, 10, 15))

    def test_conflicting_duplicate_dates_rejected(self):
        path = self.workbook([(date(2026, 10, 15), {'B': 2}), (date(2026, 10, 15), {'B': 3})])
        with self.assertRaisesRegex(ValueError, 'Conflicting duplicate'):
            read_history(path, date(2026, 10, 15))

    def test_invalid_numeric_and_nonpositive_spot_do_not_become_zero(self):
        path = self.workbook([(date(2026, 10, 15), {'V': -1, 'HW': 0, 'HX': '-2', 'HY': True})])
        rows, audit = read_history(path, date(2026, 10, 15))
        self.assertIsNone(rows[0][1]['V'])
        self.assertIsNone(rows[0][1]['HX'])
        self.assertIsNone(rows[0][1]['HY'])
        self.assertEqual(rows[0][1]['HW'], 0)
        self.assertEqual(audit['nonpositiveSpotCells'], 1)
        self.assertEqual(audit['invalidNumericCells'], 2)

    def test_private_paths_and_empty_windows_fail(self):
        with self.assertRaisesRegex(ValueError, 'outside'):
            read_history(REPO / 'private.xlsx', date(2026, 10, 15))
        with self.assertRaisesRegex(ValueError, 'start'):
            read_history(self.path, date(2026, 10, 15), date(2027, 1, 1))
        path = self.workbook([(date(2027, 1, 1), {'V': 1300})])
        with self.assertRaisesRegex(ValueError, 'No finite'):
            read_history(path, date(2026, 10, 15))

    def test_combined_snapshot_has_old_and_new_contracts_without_cma_mutation(self):
        path = self.workbook([(date(2026, 10, 15), {'V': 1300, 'HW': -1, 'HX': -2, 'HY': -3})])
        result = build(path, date(2026, 10, 16), include_model=False)
        self.assertEqual(result['kind'], 'bje_data_info')
        self.assertEqual(result['meta']['asOf'], '2026-10-16')
        self.assertEqual(result['hedgeFx']['asOf'], '2026-10-15')
        self.assertEqual(result['fxHedge']['pairs'][0]['hedge3m'], -1)
        self.assertEqual(result['hedgeFx']['pairs'][0]['history'][0]['hedgeMean'], -2)
        self.assertNotIn('assets', result)
        self.assertNotIn('fxModel', result)

    def test_model_provenance_uses_only_data_info_midpoint(self):
        path = self.workbook([(date(2026, 10, 15), {'V': 1300})])
        with patch('build_hedge_fx_snapshot.model_input', return_value=(['2026-10-15'], [1300])), \
             patch('build_hedge_fx_snapshot.build_payload', return_value={'meta': {}}) as fit:
            result = build(path, date(2026, 10, 16))
        self.assertIn('중간값', fit.call_args.kwargs['source'])
        self.assertEqual(fit.call_args.kwargs['starts'], 4)
        primary = result['fxModel']['meta']['provenance']['primary']
        self.assertEqual(primary['valueColumn'], 'V')
        self.assertFalse(result['fxModel']['meta']['provenance']['sourceSplicing'])


class TestAlignedHistories(unittest.TestCase):
    def test_all_pair_dates_match_rates_and_no_fill_from_previous_date(self):
        rows = [(date(2026, 10, 14), values(V=1300, Y=850, HW=-1, HX=-2, HY=-3, Q=3, AT=4)),
                (date(2026, 10, 15), values(V=1310, Y=None, HW=-4, HX=None, HY=-6, Q=3, AT=None))]
        result = build_histories(rows, {})
        dates = [row['date'] for row in result['usdRates']]
        for pair in result['pairs']:
            self.assertEqual([row['date'] for row in pair['history']], dates)
        jpy = next(pair for pair in result['pairs'] if pair['currency'] == 'JPYKRW')
        self.assertEqual(jpy['spotUnit'], '원/100엔')
        self.assertIsNone(jpy['history'][-1]['spot'])
        self.assertIsNone(result['usdRates'][-1]['usPr'])
        self.assertEqual(set(result['usdRates'][-1]), {'date', *RATE_FIELDS})

    def test_model_refuses_missing_on_source_day_and_retains_unchanged_quotes(self):
        rows = [(date(2026, 10, 14), values(V=1300)), (date(2026, 10, 15), values(V=1300))]
        dates, spots = model_input(rows, minimum_observations=2)
        self.assertEqual(spots, [1300, 1300])
        self.assertEqual(dates, ['2026-10-14', '2026-10-15'])
        with self.assertRaisesRegex(ValueError, 'at least'):
            model_input(rows, minimum_observations=3)
        rows[-1][1]['V'] = None
        with self.assertRaisesRegex(ValueError, 'not bridged'):
            model_input(rows, minimum_observations=2)


if __name__ == '__main__':
    unittest.main()

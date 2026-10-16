#!/usr/bin/env python3
"""Extend a private data_info snapshot with same-date Hedge/FX histories.

The workbook's hedge premium field does not specify a percentage or annualized
unit. Values and signs are retained; cross-unit differences are not precomputed.
USDKRW modelling uses only the data_info midpoint, never another quote source.
Actual workbooks and output snapshots must remain outside hahajongha/bje.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
import json
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from build_info_snapshot import (HEADERS, PAIR_SPECS, build as build_info,
                                 file_hash, outside_repo, signed_number)
from fx_markov import build_payload

RATE_FIELDS = {
    'krPr': 'Q', 'usPr': 'AT', 'kr3m': 'B', 'us3m': 'AE',
    'kr1y': 'E', 'us1y': 'AG', 'kr2y': 'G', 'us2y': 'AH',
    'kr3y': 'I', 'us3y': 'AI',
}
EXTRA_HEADERS = {
    'B': ('한국_3M', '3월이하(당일)'),
    'E': ('한국_1Y', '1년이하(당일)'),
    'G': ('한국_2Y', '2년이하(당일)'),
    'AE': ('UST3M', 'MID_Close'),
    'AG': ('UST1Y', 'MID_Close'),
    'AI': ('UST3Y', 'MID_Close'),
}
FX_UNITS = {'USDKRW': '원/달러', 'JPYKRW': '원/100엔',
            'EURKRW': '원/유로', 'AUDKRW': '원/호주달러'}
NEEDED_COLUMNS = set(RATE_FIELDS.values()) | {c for _, cs, fx in PAIR_SPECS for c in (*cs, fx)}
EXPECTED_HEADERS = {c: {**HEADERS, **EXTRA_HEADERS}[c] for c in sorted(NEEDED_COLUMNS)}
SOURCE = 'data_info.xlsx · DailyRate'
MODEL_SOURCE = 'data_info.xlsx · DailyRate!A/V · USDKRW · 중간값'


def read_history(info, as_of, start=date(2015, 1, 1)):
    """Read exact source fields; reject conflicting dates and preserve nulls."""
    if start > as_of:
        raise ValueError('History start must be on or before as-of.')
    info = outside_repo(info)
    book = load_workbook(info, read_only=True, data_only=True)
    try:
        if 'DailyRate' not in book.sheetnames:
            raise ValueError('Missing DailyRate worksheet.')
        sheet = book['DailyRate']
        header_rows = list(sheet.iter_rows(min_row=8, max_row=9, values_only=True))
        if len(header_rows) != 2 or tuple(row[0] for row in header_rows) != ('DATES', '일자'):
            raise ValueError('Expected DailyRate!A8=DATES and A9=일자.')
        indices = {}
        for column, expected in EXPECTED_HEADERS.items():
            index = column_index_from_string(column) - 1
            matches = [i for i, pair in enumerate(zip(*header_rows)) if pair == expected]
            actual = tuple(row[index] if index < len(row) else None for row in header_rows)
            if actual != expected or matches != [index]:
                raise ValueError(f'Changed or duplicate source headers at DailyRate!{column}8:9; expected {expected}.')
            indices[column] = index
        quality = {'headersVerified': len(indices), 'cachedValuesOnly': True,
                   'futureRowsExcluded': 0, 'beforeStartRowsExcluded': 0,
                   'weekendRowsExcluded': 0, 'invalidDateRows': 0,
                   'identicalDuplicateRows': 0, 'missingNumericCells': 0,
                   'invalidNumericCells': 0, 'nonpositiveSpotCells': 0,
                   'gapFilling': False, 'interpolation': False}
        spot_columns = {spec[2] for spec in PAIR_SPECS}
        dated = {}
        for row in sheet.iter_rows(min_row=10, values_only=True):
            observed = row[0]
            if isinstance(observed, datetime):
                observed = observed.date()
            if not isinstance(observed, date):
                if observed is not None:
                    quality['invalidDateRows'] += 1
                continue
            if observed > as_of:
                quality['futureRowsExcluded'] += 1
                continue
            if observed < start:
                quality['beforeStartRowsExcluded'] += 1
                continue
            if observed.weekday() > 4:
                quality['weekendRowsExcluded'] += 1
                continue
            values = {}
            for column, index in indices.items():
                raw = row[index] if index < len(row) else None
                value = signed_number(raw)
                if raw is None:
                    quality['missingNumericCells'] += 1
                elif value is None:
                    quality['invalidNumericCells'] += 1
                if column in spot_columns and value is not None and value <= 0:
                    quality['nonpositiveSpotCells'] += 1
                    value = None
                values[column] = value
            if observed in dated:
                if dated[observed] != values:
                    raise ValueError(f'Conflicting duplicate DailyRate date {observed}.')
                quality['identicalDuplicateRows'] += 1
            else:
                dated[observed] = values
        rows = [(d, dated[d]) for d in sorted(dated)]
        if not any(any(v is not None for v in values.values()) for _, values in rows):
            raise ValueError('No finite weekday observations in the requested window.')
        quality['datedRows'] = len(rows)
        quality['maxCalendarGap'] = max(((right[0] - left[0]).days for left, right in zip(rows, rows[1:])), default=0)
        return rows, quality
    finally:
        book.close()


def build_histories(rows, quality, start=date(2015, 1, 1)):
    """Keep each pair's dates aligned with rates; incomplete means stay null."""
    if not rows:
        raise ValueError('At least one source history row is required.')
    pairs = []
    pair_quality = {}
    for currency, columns, fx_column in PAIR_SPECS:
        history = []
        for observed, values in rows:
            premiums = [values.get(c) for c in columns]
            history.append({'date': observed.isoformat(), 'spot': values.get(fx_column),
                            **dict(zip(('hedge3m', 'hedge6m', 'hedge12m'), premiums)),
                            'hedgeMean': mean(premiums) if all(v is not None for v in premiums) else None})
        available = [p['date'] for p in history if p['spot'] is not None or any(p[k] is not None for k in ('hedge3m', 'hedge6m', 'hedge12m'))]
        pair_quality[currency] = {
            'missingSpotRows': sum(p['spot'] is None for p in history),
            'incompleteHedgeRows': sum(p['hedgeMean'] is None for p in history),
            'completeHedgeRows': sum(p['hedgeMean'] is not None for p in history),
            'latestSpotDate': next((p['date'] for p in reversed(history) if p['spot'] is not None), None),
            'latestCompleteHedgeDate': next((p['date'] for p in reversed(history) if p['hedgeMean'] is not None), None),
        }
        pairs.append({'currency': currency, 'spotUnit': FX_UNITS[currency],
                      'asOf': available[-1] if available else None,
                      'source': f"DailyRate!{'/'.join(columns)} · 프리미엄",
                      'spotSource': f'DailyRate!{fx_column} · {EXPECTED_HEADERS[fx_column][1]}',
                      'history': history})
    usd_rates = [{'date': d.isoformat(), **{key: values.get(c) for key, c in RATE_FIELDS.items()}}
                 for d, values in rows]
    latest = max(d for d, values in rows if any(v is not None for v in values.values()))
    return {'schemaVersion': 1, 'source': SOURCE, 'asOf': latest.isoformat(),
            'start': start.isoformat(), 'unit': 'source_value', 'unitVerified': False,
            'rateUnit': 'percent', 'pairs': pairs, 'usdRates': usd_rates,
            'rateSources': {key: f'DailyRate!{column} · {EXPECTED_HEADERS[column][0]}' for key, column in RATE_FIELDS.items()},
            'quality': {**quality, 'pairs': pair_quality,
                        'missingRateCells': sum(p[key] is None for p in usd_rates for key in RATE_FIELDS)},
            'method': [
                'Exact row 8/9 source identities; finite cached weekday observations; dates unique and sorted.',
                'All pair and rate histories use the same worksheet dates. Missing quotes remain null; no forward fill, interpolation, or zero substitution.',
                'hedgeMean requires all three same-date signed 3M/6M/12M premiums; arithmetic mean.',
                'Premium headers specify 프리미엄 only; percentage, annualization, and cost-sign conversion are not established.',
                'KR-US rates are percentage quotes. Maturity comparisons are descriptive; no cross-unit premium-minus-spread calculation is stored.',
            ]}


def model_input(rows, minimum_observations=503):
    """Do not bridge a missing/nonpositive USD quote on an observed source day."""
    selected = [(d, values.get('V')) for d, values in rows]
    if any(value is None or value <= 0 for _, value in selected):
        raise ValueError('USDKRW model requires a positive midpoint on every selected source date; missing quotes are not bridged.')
    if len(selected) < minimum_observations:
        raise ValueError(f'USDKRW model requires at least {minimum_observations} observations.')
    return [d.isoformat() for d, _ in selected], [value for _, value in selected]


def build(info, as_of, *, start=date(2015, 1, 1), include_model=True,
          horizon=504, paths=40000, seed=20261016, starts=4, max_iter=400, holdout=252):
    info = outside_repo(info)
    rows, quality = read_history(info, as_of, start)
    result = build_info(info, as_of)
    result['hedgeFx'] = build_histories(rows, quality, start)
    result['meta']['units']['hedgeFx'] = 'source_value'
    result['meta']['quality']['hedgeFx'] = quality
    if include_model:
        dates, spots = model_input(rows, minimum_observations=max(251, holdout + 251))
        model = build_payload(dates, spots, source=MODEL_SOURCE, horizon=horizon,
                              paths=paths, seed=seed, starts=starts, max_iter=max_iter, holdout=holdout)
        model['meta']['provenance'] = {
            'requestedAsOf': as_of.isoformat(), 'sampleStart': start.isoformat(),
            'primary': {'file': info.name, 'sheet': 'DailyRate', 'dateColumn': 'A',
                        'valueColumn': 'V', 'header': list(EXPECTED_HEADERS['V']), 'sha256': file_hash(info)},
            'cachedValuesOnly': True, 'gapFilling': False, 'sourceSplicing': False,
            'maxCalendarGap': quality['maxCalendarGap'],
            'transform': 'Positive cached data_info midpoint on each selected source weekday; unchanged observations retained.',
        }
        model['meta']['quoteDefinition'] = 'data_info USDKRW midpoint; no PX_LAST series used'
        result['fxModel'] = model
    json.dumps(result, ensure_ascii=False, allow_nan=False)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--info', type=Path, required=True)
    parser.add_argument('--as-of', type=date.fromisoformat, required=True)
    parser.add_argument('--start', type=date.fromisoformat, default=date(2015, 1, 1))
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--no-fx-model', action='store_true')
    parser.add_argument('--horizon', type=int, default=504)
    parser.add_argument('--paths', type=int, default=40000)
    parser.add_argument('--starts', type=int, default=4)
    parser.add_argument('--max-iter', type=int, default=400)
    args = parser.parse_args()
    try:
        output, info = outside_repo(args.output), outside_repo(args.info)
        if output == info:
            raise ValueError('Output must not overwrite the source workbook.')
        result = build(info, args.as_of, start=args.start, include_model=not args.no_fx_model,
                       horizon=args.horizon, paths=args.paths, starts=args.starts, max_iter=args.max_iter)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(',', ':')), encoding='utf-8')
    except ValueError as error:
        parser.error(str(error))
    model = result.get('fxModel')
    print(json.dumps({'kind': result['kind'], 'latestObservation': result['meta']['latestObservation'],
                      'hedgeFxRows': len(result['hedgeFx']['usdRates']), 'hedgeUnit': result['hedgeFx']['unit'],
                      'fxModelSource': model['meta']['source'] if model else None,
                      'fxModelConverged': model['model']['diagnostics']['converged'] if model else None,
                      'holdout': model['validation']['outOfSample'] if model else None}, ensure_ascii=False))


if __name__ == '__main__':
    main()

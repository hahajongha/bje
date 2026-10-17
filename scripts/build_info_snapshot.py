#!/usr/bin/env python3
"""Build a private market/KIS/hedging snapshot from cached data_info values.

Actual workbooks and generated snapshots must remain outside hahajongha/bje.
Hedge premiums retain their signed source values: the workbook does not state
an annualized unit. This builder never modifies CMA assumptions or weights.
"""
from __future__ import annotations

import argparse
import calendar
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
from statistics import mean, stdev

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

REPO = Path(__file__).resolve().parents[1]
SOURCE = 'data_info.xlsx · DailyRate'
# Exact provider row-8 / row-9 identities; positions are also verified to catch
# silent template changes before accepting a locally supplied workbook.
HEADERS = {
    'B': ('한국_3M', '3월이하(당일)'), 'E': ('한국_1Y', '1년이하(당일)'),
    'G': ('한국_2Y', '2년이하(당일)'), 'AE': ('UST3M', 'MID_Close'),
    'AG': ('UST1Y', 'MID_Close'), 'AI': ('UST3Y', 'MID_Close'),
    'Q': ('한국_기준금리', '현재가'), 'I': ('한국_3Y', '3년이하(당일)'),
    'M': ('한국_10Y', '10년이하(당일)'), 'AT': ('FFR', '현재가'),
    'AH': ('UST2Y', 'MID_Close'), 'AL': ('UST10Y', 'MID_Close'),
    'AU': ('BOJ_PR', '현재가'), 'BB': ('JPY10Y', 'MID_Close'),
    'BP': ('ECB_MRO', '현재가'), 'BV': ('GER10Y', 'MID_Close'),
    'BM': ('AUD10Y', 'MID_Close'), 'CE': ('GB10Y', 'MID_Close'),
    'IJ': ('VIX', '현재가'), 'IK': ('미국_MOVE', '현재가'),
    'II': ('VKOSPI', '현재가'), 'V': ('USDKRW', '중간값'),
    'Y': ('KRWJPY', 'KR_MID_CLOSE'), 'AB': ('EURKRW', 'KR_MID_CLOSE'),
    'AC': ('AUDKRW', 'KR_MID_CLOSE'), 'W': ('DXY', 'KR_MID_CLOSE'),
    'HW': ('USDKRW_HP_3M', '프리미엄'), 'HX': ('USDKRW_HP_6M', '프리미엄'),
    'HY': ('USDKRW_HP_12M', '프리미엄'), 'HZ': ('JPYKRW_HP_3M', '프리미엄'),
    'IA': ('JPYKRW_HP_6M', '프리미엄'), 'IB': ('JPYKRW_HP_12M', '프리미엄'),
    'IC': ('EURKRW_HP_3M', '프리미엄'), 'ID': ('EURKRW_HP_6M', '프리미엄'),
    'IE': ('EURKRW_HP_12M', '프리미엄'), 'IF': ('AUDKRW_HP_3M', '프리미엄'),
    'IG': ('AUDKRW_HP_6M', '프리미엄'), 'IH': ('AUDKRW_HP_12M', '프리미엄'),
    'JO': ('KIS채권종합', '총수익지수(산출일)'),
    'JP': ('KIS채권종합', '평균 수정 듀레이션(산출일)'),
    'JQ': ('KIS채권종합', '평균 CONVEXITY(산출일)'),
}
RATE_SPECS = [
    ('kr_policy', '한국 기준금리', 'Q'), ('kr_3y', '한국 3Y', 'I'),
    ('kr_10y', '한국 10Y', 'M'), ('us_policy', '미국 기준금리', 'AT'),
    ('us_2y', '미국 2Y', 'AH'), ('us_10y', '미국 10Y', 'AL'),
    ('jp_policy', '일본 기준금리', 'AU'), ('jp_10y', '일본 10Y', 'BB'),
    ('ecb_mro', 'ECB MRO', 'BP'), ('de_10y', '독일 10Y', 'BV'),
    ('au_10y', '호주 10Y', 'BM'), ('gb_10y', '영국 10Y', 'CE'),
]
RISK_SPECS = [('vix', 'VIX', 'IJ'), ('move', 'MOVE', 'IK'), ('vkospi', 'VKOSPI', 'II')]
FX_SPECS = [('usdkrw', 'USD/KRW', 'V', '원/달러'), ('jpykrw', 'JPY/KRW', 'Y', '원/100엔'),
            ('eurkrw', 'EUR/KRW', 'AB', '원/유로'), ('audkrw', 'AUD/KRW', 'AC', '원/호주달러'),
            ('dxy', 'DXY', 'W', '지수')]
PAIR_SPECS = [('USDKRW', ('HW', 'HX', 'HY'), 'V'), ('JPYKRW', ('HZ', 'IA', 'IB'), 'Y'),
              ('EURKRW', ('IC', 'ID', 'IE'), 'AB'), ('AUDKRW', ('IF', 'IG', 'IH'), 'AC')]
RATE_PAIR_SPECS = [('PR', 'Q', 'AT'), ('3M', 'B', 'AE'), ('1Y', 'E', 'AG'),
                   ('2Y', 'G', 'AH'), ('3Y', 'I', 'AI')]


def signed_number(value):
    """Accept zero and negatives; refuse booleans, strings and nonfinite cells."""
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return float(value)
    return None


def outside_repo(path):
    result = Path(path).resolve()
    if result.is_relative_to(REPO):
        raise ValueError('Actual source workbooks and output must be outside hahajongha/bje.')
    return result


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def read_info(path, as_of):
    path = outside_repo(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'DailyRate' not in book.sheetnames:
            raise ValueError('Missing DailyRate worksheet.')
        sheet = book['DailyRate']
        header_rows = list(sheet.iter_rows(min_row=8, max_row=9, values_only=True))
        if len(header_rows) != 2 or tuple(row[0] for row in header_rows) != ('DATES', '일자'):
            raise ValueError('Expected DailyRate!A8=DATES and A9=일자.')
        indices = {}
        for column, expected in HEADERS.items():
            index = column_index_from_string(column) - 1
            matches = [i for i, pair in enumerate(zip(*header_rows)) if pair == expected]
            actual = tuple(row[index] if index < len(row) else None for row in header_rows)
            if actual != expected or matches != [index]:
                raise ValueError(f'Changed or duplicate source headers at DailyRate!{column}8:9; expected {expected}.')
            indices[column] = matches[0]
        quality = {'futureRowsExcluded': 0, 'identicalDuplicateRows': 0, 'invalidDateRows': 0,
                   'invalidNumericCells': 0, 'missingNumericCells': 0, 'cachedValuesOnly': True,
                   'headersVerified': len(indices), 'gapFilling': False}
        dated = {}
        for row in sheet.iter_rows(min_row=10, values_only=True):
            observed = row[0]
            if isinstance(observed, datetime):
                observed = observed.date()
            if not isinstance(observed, date):
                if observed is not None:
                    quality['invalidDateRows'] += 1
                continue
            if observed > as_of:
                quality['futureRowsExcluded'] += 1
                continue
            values = {}
            for column, index in indices.items():
                raw = row[index] if index < len(row) else None
                value = signed_number(raw)
                values[column] = value
                if raw is None:
                    quality['missingNumericCells'] += 1
                elif value is None:
                    quality['invalidNumericCells'] += 1
            if observed in dated:
                if dated[observed] != values:
                    raise ValueError(f'Conflicting duplicate DailyRate date {observed}.')
                quality['identicalDuplicateRows'] += 1
            else:
                dated[observed] = values
        rows = [(d, dated[d]) for d in sorted(dated)]
        if not any(any(v is not None for v in values.values()) for _, values in rows):
            raise ValueError('No finite dated observations on or before the requested cutoff.')
        quality['datedRows'] = len(rows)
        return rows, quality
    finally:
        book.close()


def metric(rows, id_, label, column, unit):
    available = [(d, values[column]) for d, values in rows if values[column] is not None]
    if not available:
        return None
    observed, value = available[-1]
    previous_date, previous = available[-2] if len(available) > 1 else (None, None)
    return {'id': id_, 'label': label, 'value': value, 'previous': previous,
            'change': value - previous if previous is not None else None, 'unit': unit,
            'asOf': observed.isoformat(), 'previousDate': previous_date.isoformat() if previous_date else None,
            'source': f'DailyRate!{column} · {HEADERS[column][0]}'}


def build_market(rows):
    rates = [metric(rows, id_, label, column, '%') for id_, label, column in RATE_SPECS]
    risk = [metric(rows, id_, label, column, '지수') for id_, label, column in RISK_SPECS]
    fx = [metric(rows, *spec) for spec in FX_SPECS]
    return {name: [m for m in values if m is not None] for name, values in [('rates', rates), ('risk', risk), ('fx', fx)]}


def realized_volatility(rows, column):
    # Select the last 253 worksheet weekday dates BEFORE checking quotes. A
    # missing quote inside that window invalidates it rather than bridging it.
    weekday_rows = [(d, values[column]) for d, values in rows if d.weekday() < 5]
    if len(weekday_rows) < 253:
        return None, None, None
    window = weekday_rows[-253:]
    period = f'{window[0][0].isoformat()} ~ {window[-1][0].isoformat()} · 252 수익률'
    if any(v is None or v <= 0 for _, v in window):
        return None, None, period
    returns = [math.log(right[1] / left[1]) for left, right in zip(window, window[1:])]
    return stdev(returns) * math.sqrt(252) * 100, window[-1][0].isoformat(), period


def build_hedge(rows):
    pairs = []
    for currency, columns, fx_column in PAIR_SPECS:
        available = [(d, values) for d, values in rows if any(values[c] is not None for c in columns)]
        if not available:
            continue
        observed, latest = available[-1]
        history = [{'date': d.isoformat(), 'value': mean(values[c] for c in columns)}
                   for d, values in rows if d >= date(2015, 1, 1) and all(values[c] is not None for c in columns)]
        tenor_history = [
            {'date': d.isoformat(), **{key: values[c] for key, c in
                                      zip(('hedge3m', 'hedge6m', 'hedge12m'), columns)}}
            for d, values in rows if d >= date(2015, 1, 1)
            and any(values[c] is not None for c in columns)
        ]
        vol, vol_as_of, vol_period = realized_volatility(rows, fx_column)
        pairs.append({'currency': currency, 'asOf': observed.isoformat() if observed else None,
                      **{key: latest.get(c) for key, c in zip(('hedge3m', 'hedge6m', 'hedge12m'), columns)},
                      'fxVol': vol, 'volAsOf': vol_as_of, 'volPeriod': vol_period, 'aggregation': 'snapshot',
                      'source': f"DailyRate!{'/'.join(columns)}", 'history': history,
                      'tenorHistory': tenor_history})
    return {'unit': 'source_value', 'source': SOURCE, 'pairs': pairs}


def build_fx_analytics(rows, as_of):
    """Keep raw premium units and exact-date complete cases for each comparison.

    No numeric rate-minus-premium gap is produced here: source premium headers
    do not establish annualization. PR and government-curve maturities are
    comparison proxies, not maturity-matched deposit rates proving exact CIP.
    """
    history_rows = [(day, values) for day, values in rows if day >= date(2015, 1, 1)]
    rate_pairs = []
    for tenor, kr_column, us_column in RATE_PAIR_SPECS:
        required = (kr_column, us_column, 'HW', 'HX', 'HY')
        history = [
            {'date': day.isoformat(), 'kr': values[kr_column], 'us': values[us_column],
             'hedge': mean(values[column] for column in ('HW', 'HX', 'HY'))}
            for day, values in history_rows if all(values[column] is not None for column in required)
        ]
        rate_pairs.append({'tenor': tenor,
                           'krSource': f'DailyRate!{kr_column} · {HEADERS[kr_column][0]}',
                           'usSource': f'DailyRate!{us_column} · {HEADERS[us_column][0]}',
                           'history': history})
    spots = [
        {'id': identifier, 'label': label, 'unit': unit,
         'source': f'DailyRate!{column} · {HEADERS[column][0]} · {HEADERS[column][1]}',
         'history': [{'date': day.isoformat(), 'value': values[column]}
                     for day, values in history_rows if values[column] is not None and values[column] > 0]}
        for identifier, label, column, unit in FX_SPECS
    ]
    return {'schemaVersion': 1, 'asOf': as_of.isoformat(), 'source': SOURCE,
            'hedgeUnit': 'source_value', 'ratePairs': rate_pairs, 'spots': spots}


def month_shift(year, month, offset):
    index = year * 12 + month - 1 + offset
    return index // 12, index % 12 + 1


def completed_month_endpoints(common, as_of, months=60):
    if months < 2:
        raise ValueError('At least two monthly returns are required.')
    final_year, final_month = month_shift(as_of.year, as_of.month, -1)
    endpoints = []
    for offset in range(-months, 1):
        year, month = month_shift(final_year, final_month, offset)
        candidates = [d for d in common if d.year == year and d.month == month]
        if not candidates:
            raise ValueError(f'Missing completed month {year}-{month:02d}; no return bridging.')
        endpoint = max(candidates)
        if calendar.monthrange(year, month)[1] - endpoint.day > 7:
            raise ValueError(f'Stale month-end endpoint {endpoint}; maximum calendar lag is 7 days.')
        endpoints.append(endpoint)
    return endpoints


def build_bond(rows, as_of, months=60):
    common = {d: values for d, values in rows if all(values[c] is not None and values[c] > 0 for c in ('JO', 'JP', 'JQ'))}
    index = {d: values['JO'] for d, values in rows if values['JO'] is not None and values['JO'] > 0}
    result = {'name': 'KIS 종합채권', 'asOf': None, 'index': None, 'duration': None, 'convexity': None,
              'totalReturnYtd': None, 'totalReturn1y': None, 'annualMean': None, 'annualVol': None,
              'annualGeometricReturn': None, 'maxDrawdown': None, 'periodStart': None, 'periodEnd': None,
              'observations': 0, 'monthly': [], 'source': 'DailyRate!JO: JQ'}
    if not common:
        return result, 'No positive common index/duration/convexity observation.'
    latest = max(common)
    values = common[latest]
    result.update(asOf=latest.isoformat(), index=values['JO'], duration=values['JP'], convexity=values['JQ'])
    prior_year_end = date(latest.year - 1, 12, 31)
    year_base = [d for d in index if d <= prior_year_end and (prior_year_end - d).days <= 7]
    if year_base:
        result['totalReturnYtd'] = (values['JO'] / index[max(year_base)] - 1) * 100
    try:
        anniversary = latest.replace(year=latest.year - 1)
    except ValueError:
        anniversary = latest.replace(year=latest.year - 1, day=28)
    year_ago = [d for d in index if d <= anniversary and (anniversary - d).days <= 7]
    if year_ago:
        result['totalReturn1y'] = (values['JO'] / index[max(year_ago)] - 1) * 100
    try:
        endpoints = completed_month_endpoints(common, as_of, months)
    except ValueError as error:
        return result, str(error)
    levels = [common[d]['JO'] for d in endpoints]
    returns = [right / left - 1 for left, right in zip(levels, levels[1:])]
    peak, maximum_drawdown = levels[0], 0.0
    for level in levels:
        peak = max(peak, level)
        maximum_drawdown = min(maximum_drawdown, level / peak - 1)
    result.update(annualMean=mean(returns) * 12 * 100, annualVol=stdev(returns) * math.sqrt(12) * 100,
                  annualGeometricReturn=((levels[-1] / levels[0]) ** (12 / len(returns)) - 1) * 100,
                  maxDrawdown=maximum_drawdown * 100, periodStart=endpoints[0].isoformat(),
                  periodEnd=endpoints[-1].isoformat(), observations=len(returns),
                  monthly=[{'date': d.isoformat(), 'index': common[d]['JO']} for d in endpoints])
    return result, None


def build(info, as_of):
    info = outside_repo(info)
    rows, quality = read_info(info, as_of)
    return build_from_rows(info, as_of, rows, quality)


def build_from_rows(info, as_of, rows, quality):
    """Share one audited workbook read with the optional combined FX builder."""
    info = outside_repo(info)
    quality = dict(quality)
    bond, bond_issue = build_bond(rows, as_of)
    quality['bondMonthlyIssue'] = bond_issue
    if bond['asOf'] is None:
        bond = None
    hedge = build_hedge(rows)
    present_currencies = {pair['currency'] for pair in hedge['pairs']}
    quality['missingHedgeCurrencies'] = [currency for currency, _, _ in PAIR_SPECS if currency not in present_currencies]
    latest = max(d for d, values in rows if any(v is not None for v in values.values()))
    result = {'kind': 'bje_data_info', 'schemaVersion': 1,
              'meta': {'source': SOURCE, 'asOf': as_of.isoformat(), 'latestObservation': latest.isoformat(),
                       'sha256': file_hash(info), 'units': {'hedge': 'source_value', 'rates': 'percent',
                       'bondReturns': 'percent', 'fxVolatility': 'annual_percent'}, 'quality': quality,
                       'method': ['Cached finite signed values; exact row8/9 header verification; no filling or interpolation.',
                                  'Metric changes are latest minus previous available observation; dates are retained.',
                                  'Hedge premium units are unspecified in the source headers; signed source values are unchanged.',
                                  'Hedge history: complete same-date 3M/6M/12M arithmetic averages since 2015-01-01.',
                                  'Hedge tenorHistory: signed individual 3M/6M/12M premiums since 2015-01-01; partial dates retain nulls without filling.',
                                  'Hedge rate comparisons: same-date complete KR/US rates and HW/HX/HY mean; raw source premium units retained, no inferred CIP gap.',
                                  'FX monitoring: positive finite observed levels since 2015-01-01; source timestamps retained without filling or source blending.',
                                  'FX volatility: 252 log returns from the last 253 worksheet weekday dates; sample std × sqrt(252) × 100. Missing/nonpositive window quotes invalidate volatility.',
                                  'KIS: 60 monthly returns through the previous calendar month, 61 common index/duration/convexity endpoints within 7 calendar days of month end.',
                                  'KIS annual mean: monthly arithmetic mean × 12 × 100; sample monthly std × sqrt(12) × 100; geometric return uses actual month count; drawdown uses monthly endpoints.',
                                  'KIS YTD uses the final prior-year index within 7 calendar days of Dec 31; 1Y uses the last index at/before anniversary with at most 7 calendar days lag.']},
              'market': build_market(rows), 'bond': bond, 'fxHedge': hedge,
              'fxAnalytics': build_fx_analytics(rows, as_of)}
    # Validate JSON finiteness before any write.
    json.dumps(result, ensure_ascii=False, allow_nan=False)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--info', type=Path, required=True)
    parser.add_argument('--as-of', type=date.fromisoformat, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        output = outside_repo(args.output)
        info = outside_repo(args.info)
        if output == info:
            raise ValueError('Output must not overwrite the source workbook.')
        result = build(info, args.as_of)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(',', ':')), encoding='utf-8')
    except ValueError as error:
        parser.error(str(error))
    print(json.dumps({'kind': result['kind'], 'asOf': result['meta']['asOf'],
                      'latestObservation': result['meta']['latestObservation'],
                      'bondMonths': result['bond']['observations'] if result['bond'] else 0,
                      'hedgeHistory': {p['currency']: len(p['history']) for p in result['fxHedge']['pairs']},
                      'quality': result['meta']['quality']}, ensure_ascii=False))


if __name__ == '__main__':
    main()

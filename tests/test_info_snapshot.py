"""Synthetic-only tests for the private data_info snapshot contract."""
import calendar
from datetime import date, timedelta
import math
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from build_info_snapshot import (HEADERS, REPO, build_bond, build_hedge, build_market,
                                 completed_month_endpoints, outside_repo, read_info,
                                 realized_volatility, signed_number)


def values(**updates):
    return {**dict.fromkeys(HEADERS), **updates}


def weekdays(start, n):
    result = []
    while len(result) < n:
        if start.weekday() < 5:
            result.append(start)
        start += timedelta(days=1)
    return result


class TestInfoReader(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'synthetic.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows, transform=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'DailyRate'
        sheet.cell(8, 1, 'DATES')
        sheet.cell(9, 1, '일자')
        for column, labels in HEADERS.items():
            index = column_index_from_string(column)
            sheet.cell(8, index, labels[0])
            sheet.cell(9, index, labels[1])
        for n, (observed, data) in enumerate(rows, 10):
            sheet.cell(n, 1, observed)
            for column, value in data.items():
                sheet.cell(n, column_index_from_string(column), value)
        if transform:
            transform(sheet)
        book.save(self.path)
        book.close()
        return self.path

    def test_signed_values_sort_cutoff_missing_and_duplicate(self):
        path = self.workbook([(date(2026, 10, 15), {'HW': -1.25, 'AU': -0.1}),
                              (date(2026, 10, 14), {'HW': 0, 'AU': -0.2}),
                              (date(2026, 10, 15), {'HW': -1.25, 'AU': -0.1}),
                              (date(2026, 10, 16), {'HW': 10000})])
        rows, audit = read_info(path, date(2026, 10, 15))
        self.assertEqual([d for d, _ in rows], [date(2026, 10, 14), date(2026, 10, 15)])
        self.assertEqual([v['HW'] for _, v in rows], [0, -1.25])
        self.assertIsNone(rows[-1][1]['HX'])
        self.assertEqual(audit['futureRowsExcluded'], 1)
        self.assertEqual(audit['identicalDuplicateRows'], 1)
        self.assertEqual(audit['headersVerified'], len(HEADERS))
        market = build_market(rows)
        self.assertEqual(len(market['rates']), 1)
        self.assertEqual(market['rates'][0]['value'], -0.1)
        self.assertAlmostEqual(market['rates'][0]['change'], 0.1)
        self.assertEqual(market['rates'][0]['unit'], '%')
        self.assertEqual(market['fx'], [])

    def test_conflicting_duplicate_even_with_missing_quote(self):
        for second in [2, None]:
            path = self.workbook([(date(2026, 1, 2), {'HW': 1}), (date(2026, 1, 2), {'HW': second})])
            with self.assertRaisesRegex(ValueError, 'Conflicting duplicate'):
                read_info(path, date(2026, 1, 2))

    def test_exact_both_header_rows_and_duplicate_identity(self):
        for transform in [lambda s: s.__setitem__('HW9', 'annual_percent'),
                          lambda s: s.__setitem__('HW8', 'JPYKRW_HP_3M'),
                          lambda s: s.__setitem__('A8', 'Dates'),
                          lambda s: (s.__setitem__('B8', 'USDKRW_HP_3M'), s.__setitem__('B9', '프리미엄'))]:
            path = self.workbook([(date(2026, 1, 2), {'HW': -1})], transform)
            with self.assertRaises(ValueError):
                read_info(path, date(2026, 1, 2))

    def test_moved_and_absent_headers(self):
        def move_hw_and_drop_move(sheet):
            sheet.insert_cols(2)
            sheet['B8'], sheet['B9'] = 'USDKRW_HP_3M', '프리미엄'
            sheet.delete_cols(column_index_from_string('HX'))
            # Columns before the old HW shift right by one; later ones are unchanged.
            for column in ('IK', 'JO', 'JP', 'JQ'):
                sheet[column + '8'] = sheet[column + '9'] = None
            for n in range(10, sheet.max_row + 1):
                sheet[f'B{n}'] = -7
        path = self.workbook([(date(2026, 1, 2), {'HW': -1, 'IK': 90, 'JO': 100})], move_hw_and_drop_move)
        rows, audit = read_info(path, date(2026, 1, 2))
        self.assertEqual(rows[0][1]['HW'], -7)
        self.assertIsNone(rows[0][1]['IK'])
        self.assertIsNone(rows[0][1]['JO'])
        self.assertEqual(audit['missingHeaders'], ['미국_MOVE', 'KIS채권종합 총수익지수(산출일)',
                                                   'KIS채권종합 평균 수정 듀레이션(산출일)',
                                                   'KIS채권종합 평균 CONVEXITY(산출일)'])
        self.assertEqual(audit['headersVerified'], len(HEADERS) - 4)
        self.assertEqual(build_hedge(rows)['pairs'][0]['source'], 'DailyRate!B/HX/HY')

    def test_numeric_type_gate_and_repo_safety(self):
        self.assertEqual(signed_number(-2), -2.0)
        self.assertEqual(signed_number(0), 0.0)
        for raw in [float('nan'), float('inf'), -float('inf'), True, '-2', '', None]:
            self.assertIsNone(signed_number(raw))
        with self.assertRaisesRegex(ValueError, 'outside'):
            outside_repo(REPO / 'data' / 'private.json')
        with self.assertRaisesRegex(ValueError, 'outside'):
            outside_repo(REPO / 'scripts' / '..' / 'source.xlsx')
        self.assertEqual(outside_repo(self.path), self.path.resolve())


class TestHedgeSnapshot(unittest.TestCase):
    def test_latest_partial_keeps_missing_and_history_complete_only(self):
        rows = [(date(2014, 12, 31), values(HW=-99, HX=-99, HY=-99)),
                (date(2015, 1, 2), values(HW=-3, HX=0, HY=3)),
                (date(2026, 10, 14), values(HW=-1, HX=-2, HY=-3)),
                (date(2026, 10, 15), values(HW=-4, HX=None, HY=-5))]
        result = build_hedge(rows)
        usd = result['pairs'][0]
        self.assertEqual(result['unit'], 'source_value')
        self.assertEqual(usd['asOf'], '2026-10-15')
        self.assertEqual(usd['hedge3m'], -4)
        self.assertIsNone(usd['hedge6m'])
        self.assertEqual(usd['history'], [{'date': '2015-01-02', 'value': 0.0}, {'date': '2026-10-14', 'value': -2.0}])
        self.assertEqual(usd['source'], 'DailyRate!HW/HX/HY')
        self.assertEqual(len(result['pairs']), 1)
        self.assertEqual(usd['currency'], 'USDKRW')

    def test_volatility_against_known_log_returns_and_253_boundary(self):
        dates = weekdays(date(2025, 1, 1), 253)
        levels = [100.0]
        for k in range(252):
            levels.append(levels[-1] * math.exp(0.01 if k % 2 else -0.01))
        rows = [(d, values(V=v)) for d, v in zip(dates, levels)]
        vol, observed, period = realized_volatility(rows, 'V')
        expected = math.sqrt(252 * 0.01**2 / 251) * math.sqrt(252) * 100
        self.assertAlmostEqual(vol, expected, places=10)
        self.assertEqual(observed, dates[-1].isoformat())
        self.assertIn('252', period)
        self.assertIsNone(realized_volatility(rows[:-1], 'V')[0])
        for invalid in [None, 0, -1]:
            damaged = list(rows)
            damaged[-2] = (damaged[-2][0], values(V=invalid))
            self.assertIsNone(realized_volatility(damaged, 'V')[0])
        # A missing quote outside the last 253 dates does not invalidate it.
        self.assertAlmostEqual(realized_volatility([(date(2024, 12, 31), values())] + rows, 'V')[0], expected)

    def test_missing_latest_does_not_backfill_volatility(self):
        rows = [(d, values(V=100)) for d in weekdays(date(2025, 1, 1), 254)]
        rows[-1] = (rows[-1][0], values(V=None))
        vol, observed, period = realized_volatility(rows, 'V')
        self.assertIsNone(vol)
        self.assertIsNone(observed)
        self.assertIn(rows[-1][0].isoformat(), period)

    def test_jpy_quote_retains_100_yen_unit(self):
        result = build_market([(date(2026, 10, 15), values(Y=851.5))])
        self.assertEqual(result['fx'][0]['unit'], '원/100엔')
        self.assertEqual(result['fx'][0]['value'], 851.5)


class TestBondSnapshot(unittest.TestCase):
    @staticmethod
    def monthly_rows(months=60):
        rows = []
        level = 100.0
        for k in range(months + 1):
            index = 2021 * 12 + 8 + k
            year, month = index // 12, index % 12 + 1
            d = date(year, month, calendar.monthrange(year, month)[1])
            if k:
                level *= 1.01
            rows.append((d, values(JO=level, JP=5, JQ=80)))
        return rows

    def test_60_complete_months_and_constant_return_identity(self):
        rows = self.monthly_rows()
        rows.append((date(2026, 10, 15), values(JO=1000, JP=6, JQ=90)))
        result, issue = build_bond(rows, date(2026, 10, 16))
        self.assertIsNone(issue)
        self.assertEqual(result['observations'], 60)
        self.assertEqual(len(result['monthly']), 61)
        self.assertEqual(result['periodStart'], '2021-09-30')
        self.assertEqual(result['periodEnd'], '2026-09-30')
        self.assertEqual(result['asOf'], '2026-10-15')
        self.assertEqual(result['index'], 1000)
        self.assertAlmostEqual(result['annualMean'], 12)
        self.assertAlmostEqual(result['annualVol'], 0)
        self.assertAlmostEqual(result['annualGeometricReturn'], (1.01**12 - 1) * 100)
        self.assertEqual(result['maxDrawdown'], 0)

    def test_month_end_seven_day_boundary_and_no_missing_month_bridge(self):
        common = {d: v for d, v in self.monthly_rows()}
        end = date(2026, 9, 30)
        v = common.pop(end)
        common[end - timedelta(days=7)] = v
        self.assertEqual(completed_month_endpoints(common, date(2026, 10, 16))[-1], date(2026, 9, 23))
        common[end - timedelta(days=8)] = common.pop(end - timedelta(days=7))
        with self.assertRaisesRegex(ValueError, 'Stale'):
            completed_month_endpoints(common, date(2026, 10, 16))
        del common[end - timedelta(days=8)]
        with self.assertRaisesRegex(ValueError, 'Missing completed month'):
            completed_month_endpoints(common, date(2026, 10, 16))

    def test_missing_duration_removes_common_month_not_other_months(self):
        rows = self.monthly_rows()
        rows[20][1]['JP'] = None
        result, issue = build_bond(rows, date(2026, 10, 16))
        self.assertIn('Missing completed month', issue)
        self.assertEqual(result['monthly'], [])
        self.assertEqual(result['observations'], 0)
        self.assertIsNone(result['annualMean'])
        self.assertIsNotNone(result['index'])

    def test_ytd_anniversary_lag_and_monthly_drawdown(self):
        rows = [(date(2025, 10, 8), values(JO=100, JP=5, JQ=80)),
                (date(2025, 12, 31), values(JO=110, JP=5, JQ=80)),
                (date(2026, 10, 15), values(JO=121, JP=5, JQ=80))]
        result, _ = build_bond(rows, date(2026, 10, 16))
        self.assertAlmostEqual(result['totalReturnYtd'], 10)
        self.assertAlmostEqual(result['totalReturn1y'], 21)
        rows[0] = (date(2025, 10, 7), rows[0][1])
        result, _ = build_bond(rows, date(2026, 10, 16))
        self.assertIsNone(result['totalReturn1y'])
        rows[1] = (date(2025, 12, 23), rows[1][1])
        result, _ = build_bond(rows, date(2026, 10, 16))
        self.assertIsNone(result['totalReturnYtd'])
        tiny = [(date(2026, 7, 31), values(JO=100, JP=5, JQ=80)),
                (date(2026, 8, 31), values(JO=80, JP=5, JQ=80)),
                (date(2026, 9, 30), values(JO=90, JP=5, JQ=80))]
        result, issue = build_bond(tiny, date(2026, 10, 16), months=2)
        self.assertIsNone(issue)
        self.assertAlmostEqual(result['maxDrawdown'], -20)
        self.assertAlmostEqual(result['annualMean'], (-0.2 + 0.125) / 2 * 1200)
        self.assertAlmostEqual(result['annualVol'], math.sqrt((0.325**2) / 2) * math.sqrt(12) * 100)
        self.assertAlmostEqual(result['annualGeometricReturn'], (0.9**6 - 1) * 100)


if __name__ == '__main__':
    unittest.main()

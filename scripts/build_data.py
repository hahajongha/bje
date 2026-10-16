#!/usr/bin/env python3
"""Build a private, KRW total-return dataset from user-supplied workbooks.

The source files and output must stay outside the public repository. This
script reads cached workbook values; it does not refresh data-provider add-ins.
Example:
  python scripts/build_data.py --info /private/data_info.xlsx \
    --bb /private/data_bb.xlsx --as-of 2026-10-15 --output /private/portfolio.json
"""

from __future__ import annotations

import argparse
import calendar
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
from zoneinfo import ZoneInfo

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


REPO = Path(__file__).resolve().parents[1]
SPECS = [
    dict(id="kr_bond", name="국내 종합채권", ticker="KIS채권종합", source="info", column="JO", color="#55d5ce", currentWeight=0.35),
    dict(id="kr_equity", name="국내 주식", ticker="KOSPI TR", source="bb", column="FX", color="#94a8ff", currentWeight=0.15),
    dict(id="us_equity", name="미국 주식", ticker="S&P 500 TR", source="bb", column="FW", color="#f4c975", currentWeight=0.20),
    dict(id="us_treasury", name="미국 국채", ticker="IEF TR", source="bb", column="GP", color="#7db7ed", currentWeight=0.15),
    dict(id="us_credit", name="미국 투자등급", ticker="LQD TR", source="bb", column="GR", color="#c9a2ec", currentWeight=0.10),
    dict(id="usd_bills", name="달러 단기채", ticker="BIL TR", source="bb", column="GT", color="#e19c8b", currentWeight=0.05),
]


def valid_level(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_series(path, sheet, columns, as_of):
    book = load_workbook(path, read_only=True, data_only=True)
    if sheet not in book.sheetnames:
        raise ValueError(f"Missing source worksheet: {sheet}")
    ws = book[sheet]
    headers = list(ws.iter_rows(min_row=6, max_row=9, values_only=True))
    indexes = {col: column_index_from_string(col) - 1 for col in columns}
    series = {col: {} for col in columns}
    for row in ws.iter_rows(min_row=10, values_only=True):
        observed = row[0]
        if isinstance(observed, datetime):
            observed = observed.date()
        if not isinstance(observed, date) or observed > as_of:
            continue
        for col, index in indexes.items():
            value = row[index] if index < len(row) else None
            if valid_level(value):
                if observed in series[col] and series[col][observed] != float(value):
                    raise ValueError(f"Conflicting duplicate {sheet}!{col} at {observed}")
                series[col][observed] = float(value)
    book.close()
    if any(not values for values in series.values()):
        raise ValueError(f"Empty positive numeric series: {path.name}")
    labels = {col: [str(row[index] or "") for row in headers] for col, index in indexes.items()}
    return series, labels


def source_currencies(path, columns):
    """Read provider requests so foreign ETF values are not converted twice."""
    book = load_workbook(path, read_only=True, data_only=False)
    row = next(book["D"].iter_rows(min_row=10, max_row=10, values_only=True))
    currencies = {}
    for col in columns:
        request = str(row[column_index_from_string(col) - 1])
        match = re.search(r'CURR=([A-Z]{3})', request, re.I)
        if not match:
            raise ValueError(f"No explicit source currency in D!{col}10")
        currencies[col] = match.group(1).upper()
    book.close()
    return currencies


def shift_month(year, month, offset):
    index = year * 12 + month - 1 + offset
    return index // 12, index % 12 + 1


def build(info_path, bb_path, as_of, months):
    info, info_headers = read_series(info_path, "DailyRate", ["JO"], as_of)
    bb_cols = ["BB", "FW", "FX", "GP", "GR", "GT"]
    bb, bb_headers = read_series(bb_path, "D", bb_cols, as_of)
    currencies = source_currencies(bb_path, bb_cols)
    if "총수익지수" not in info_headers["JO"][3] or "KIS채권종합" not in info_headers["JO"]:
        raise ValueError("DailyRate!JO no longer identifies the KIS total-return index")
    expected_tickers = {"BB": "USDKRW CURNCY", "FW": "SPX INDEX", "FX": "KOSPI INDEX", "GP": "IEF US EQUITY", "GR": "LQD US EQUITY", "GT": "BIL US EQUITY"}
    for col, ticker in expected_tickers.items():
        if bb_headers[col][2].strip().upper() != ticker:
            raise ValueError(f"Source mapping changed at D!{col}8")
        if col != "BB" and bb_headers[col][3] != "TOT_RETURN_INDEX_GROSS_DVDS":
            raise ValueError(f"D!{col} no longer requests total returns")
    if currencies["FW"] != "USD" or any(currencies[c] != "KRW" for c in ["FX", "GP", "GR", "GT"]):
        raise ValueError("Source currency assumptions changed; inspect before converting")

    # Use matching dated observations. No carry-forward, interpolation, or
    # nonconsecutive monthly return is introduced by this extractor.
    raw = [info["JO"], bb["FX"], bb["FW"], bb["GP"], bb["GR"], bb["GT"], bb["BB"]]
    common_dates = set.intersection(*(set(series) for series in raw))
    if not common_dates:
        raise ValueError("No common dated observations across all inputs")
    end_year, end_month = shift_month(as_of.year, as_of.month, -1)
    month_keys = [shift_month(end_year, end_month, offset) for offset in range(-months, 1)]
    endpoints = []
    endpoint_audit = []
    for year, month in month_keys:
        candidates = [d for d in common_dates if d.year == year and d.month == month]
        if not candidates:
            raise ValueError(f"Missing complete month: {year}-{month:02d}; no multi-month bridging allowed")
        endpoint = max(candidates)
        stale_days = calendar.monthrange(year, month)[1] - endpoint.day
        if stale_days > 7:
            raise ValueError(f"Month-end endpoint is stale by {stale_days} days: {endpoint}")
        endpoints.append(endpoint)
        endpoint_audit.append({"month": f"{year}-{month:02d}", "date": endpoint.isoformat(),
                               "calendarDayLag": stale_days, "commonDailyObservations": len(candidates)})

    levels = np.array([
        [info["JO"][d], bb["FX"][d], bb["FW"][d] * bb["BB"][d], bb["GP"][d], bb["GR"][d], bb["GT"][d]]
        for d in endpoints
    ], dtype=float)
    returns = levels[1:] / levels[:-1] - 1.0
    means = returns.mean(axis=0) * 12.0
    covariance = np.cov(returns, rowvar=False, ddof=1) * 12.0
    volatility = np.sqrt(np.diag(covariance))
    correlation = np.corrcoef(returns, rowvar=False)
    base100 = levels / levels[0] * 100.0
    geometric = (levels[-1] / levels[0]) ** (12.0 / months) - 1.0
    drawdown = base100 / np.maximum.accumulate(base100, axis=0) - 1.0
    if not np.isfinite(returns).all() or not np.isfinite(covariance).all():
        raise ValueError("Non-finite calculated return or covariance")
    if np.linalg.eigvalsh(covariance).min() < -1e-12:
        raise ValueError("The complete-case covariance is not positive semidefinite")
    if not np.allclose(covariance, correlation * np.outer(volatility, volatility), atol=1e-12):
        raise ValueError("Covariance / correlation reconciliation failed")

    assets = []
    for i, spec in enumerate(SPECS):
        is_info = spec["source"] == "info"
        source = {
            "file": info_path.name if is_info else bb_path.name,
            "sheet": "DailyRate" if is_info else "D",
            "column": spec["column"],
            "field": "총수익지수(산출일)" if is_info else "TOT_RETURN_INDEX_GROSS_DVDS",
            "inputCurrency": "USD" if spec["id"] == "us_equity" else "KRW",
            "conversion": "동일 날짜 S&P 500 USD 총수익지수 × USDKRW" if spec["id"] == "us_equity" else "원화 입력값 사용; 추가 환산 없음",
        }
        assets.append({
            **{key: spec[key] for key in ["id", "name", "ticker", "color", "currentWeight"]},
            "currency": "KRW", "returnType": "total_return", "source": source,
            "annualMean": float(means[i]), "annualVol": float(volatility[i]),
            "annualGeometricReturn": float(geometric[i]), "maxDrawdown": float(drawdown[:, i].min()),
        })

    return {
        "schemaVersion": 1,
        "meta": {
            "title": "KRW 자산배분 입력 데이터", "mode": "private_actual",
            "asOf": as_of.isoformat(), "baseCurrency": "KRW", "frequency": "monthly",
            "units": {"returns": "decimal", "weights": "decimal", "covariance": "decimal_squared_per_year", "volatility": "decimal_per_sqrt_year"},
            "periodStart": endpoints[0].isoformat(), "periodEnd": endpoints[-1].isoformat(),
            "returnStart": endpoints[1].isoformat(), "observations": months,
            "latestCommonObservation": max(common_dates).isoformat(),
            "weightStatus": "illustrative", "forecastStatus": "historical_reference_only",
            "files": [{"name": p.name, "sha256": file_hash(p)} for p in [info_path, bb_path]],
            "sourceCoverage": [
                {"file": info_path.name, "column": "JO", "first": min(info["JO"]).isoformat(), "last": max(info["JO"]).isoformat()},
                *[{"file": bb_path.name, "column": c, "first": min(bb[c]).isoformat(), "last": max(bb[c]).isoformat()} for c in bb_cols],
            ],
            "dataQuality": {"missingMonthlyEndpoints": 0, "interpolatedObservations": 0,
                            "maxCalendarDayLag": max(a["calendarDayLag"] for a in endpoint_audit),
                            "monthlyEndpoints": endpoint_audit},
            "method": "동일 날짜 공통 관측의 월별 마지막 값; 단순 월수익률. 산술평균 × 12, 표본공분산(ddof=1) × 12, 표준편차 × sqrt(12).",
            "notes": [
                "모든 자산은 원화 총수익 기준이며 해외자산은 비헤지 노출입니다.",
                "미국 ETF의 원자료는 CURR=KRW 요청값입니다. S&P 500 USD 총수익지수만 USDKRW로 환산합니다.",
                "현재 비중은 예시이며 실제 보유현황이 아닙니다. 과거 산술평균은 전망이나 추천 기대수익률이 아닙니다.",
                "월수익률은 달력상 연속된 완결월만 사용합니다. 현재 진행 중인 달은 제외합니다.",
                "월말은 모든 입력에 값이 있는 마지막 공통 날짜입니다. 거래시간 차이·원자료의 휴장일 반복값은 남을 수 있습니다.",
                "지수와 ETF 총수익은 구현·비용 구조가 다릅니다. 세금·거래비용·실제 리밸런싱 비용은 반영하지 않습니다.",
                "원본 통합문서의 저장된 값과 수식 메타데이터를 확인했으며 데이터 제공기관의 독립 원자료와 대조하지 않았습니다.",
            ],
        },
        "assets": assets,
        "covariance": covariance.tolist(), "correlation": correlation.tolist(),
        "monthly": {"dates": [d.isoformat() for d in endpoints[1:]], "returns": returns.tolist(),
                    "baseDates": [d.isoformat() for d in endpoints], "base100": base100.tolist()},
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--info", type=Path, required=True)
    parser.add_argument("--bb", type=Path, required=True)
    parser.add_argument("--as-of", type=date.fromisoformat, default=datetime.now(ZoneInfo("Asia/Seoul")).date())
    parser.add_argument("--months", type=int, default=60)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.months < 12:
        parser.error("Use at least 12 monthly returns")
    for path in [args.info, args.bb, args.output]:
        if path.resolve().is_relative_to(REPO):
            parser.error("Private source workbooks and output must be outside the public repository")
    result = build(args.info, args.bb, args.as_of, args.months)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "assets": len(result["assets"]), "monthlyReturns": result["meta"]["observations"],
                      "periodStart": result["meta"]["periodStart"], "periodEnd": result["meta"]["periodEnd"], "latestCommonObservation": result["meta"]["latestCommonObservation"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
